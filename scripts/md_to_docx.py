"""Convert FULL_PROJECT_DOCUMENTATION.md to a Word .docx file."""
from __future__ import annotations

import re
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "FULL_PROJECT_DOCUMENTATION.md"
OUT = ROOT / "docs" / "RFH_HMS_Full_Project_Documentation.docx"


def set_run_font(run, size=11, bold=False, italic=False, code=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = "Consolas" if code else "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Calibri")
    if code:
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def add_inline(paragraph, node):
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            run = paragraph.add_run(text)
            set_run_font(run)
        return
    if not isinstance(node, Tag):
        return
    name = node.name.lower()
    if name in ("strong", "b"):
        run = paragraph.add_run(node.get_text())
        set_run_font(run, bold=True)
    elif name in ("em", "i"):
        run = paragraph.add_run(node.get_text())
        set_run_font(run, italic=True)
    elif name == "code":
        run = paragraph.add_run(node.get_text())
        set_run_font(run, size=10, code=True)
    elif name == "a":
        run = paragraph.add_run(node.get_text())
        set_run_font(run)
        run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        run.underline = True
    elif name == "br":
        paragraph.add_run().add_break()
    else:
        for child in node.children:
            add_inline(paragraph, child)


def add_paragraph_from_tag(doc, tag, style=None):
    p = doc.add_paragraph(style=style)
    for child in tag.children:
        add_inline(p, child)
    return p


def add_table(doc, table_tag):
    rows = table_tag.find_all("tr")
    if not rows:
        return
    cols = max(len(r.find_all(["th", "td"])) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        cells = row.find_all(["th", "td"])
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            if j < len(cells):
                for child in cells[j].children:
                    add_inline(p, child)
                if cells[j].name == "th":
                    for run in p.runs:
                        run.bold = True
    doc.add_paragraph()


def convert():
    md_text = SRC.read_text(encoding="utf-8")
    # Keep mermaid fences as labeled code blocks for Word
    md_text = re.sub(
        r"```mermaid\n([\s\S]*?)```",
        lambda m: "```\n[Mermaid diagram — paste into mermaid.live to render]\n"
        + m.group(1).rstrip()
        + "\n```",
        md_text,
    )

    html = markdown.markdown(
        md_text,
        extensions=["tables", "fenced_code", "sane_lists", "toc"],
    )
    soup = BeautifulSoup(html, "html.parser")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    # Title page block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Rob Ferreira Hospital Management System (RFH HMS)")
    set_run_font(r, size=22, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Full Project Documentation")
    set_run_font(r, size=16, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Project Plan · IEEE 830 SRS · Agile/Scrum · UML Diagrams · Scholarly References\n"
        "Version 1.0 · September 2026"
    )
    set_run_font(r, size=11, italic=True)
    doc.add_page_break()

    for el in soup.children:
        if isinstance(el, NavigableString):
            continue
        if not isinstance(el, Tag):
            continue
        name = el.name.lower()
        if name == "h1":
            p = add_paragraph_from_tag(doc, el, style="Heading 1")
        elif name == "h2":
            p = add_paragraph_from_tag(doc, el, style="Heading 2")
        elif name == "h3":
            p = add_paragraph_from_tag(doc, el, style="Heading 3")
        elif name == "h4":
            p = add_paragraph_from_tag(doc, el, style="Heading 4")
        elif name == "p":
            add_paragraph_from_tag(doc, el)
        elif name == "ul":
            for li in el.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Bullet")
                for child in li.children:
                    add_inline(p, child)
        elif name == "ol":
            for li in el.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Number")
                for child in li.children:
                    add_inline(p, child)
        elif name == "pre":
            code = el.get_text()
            p = doc.add_paragraph()
            run = p.add_run(code)
            set_run_font(run, size=9, code=True)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Inches(0.15)
        elif name == "blockquote":
            p = add_paragraph_from_tag(doc, el)
            p.paragraph_format.left_indent = Inches(0.3)
            for run in p.runs:
                run.italic = True
        elif name == "table":
            add_table(doc, el)
        elif name == "hr":
            doc.add_paragraph("─" * 40)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    convert()
